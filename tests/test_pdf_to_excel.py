"""
Tests for pdf_to_excel.py.

Covers:
- Amount parsing
- Journal table parsing (headers, separators, page breaks, totals, other sections)
- Excel export
- Conversion of the sample PDF in docs/
"""

import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal

from pdf_to_excel import COLUMNS, main, parse_amount, parse_journal_lines, write_excel

SAMPLE_PDF = os.path.join(os.path.dirname(__file__), '..', 'docs', '2025.05.FA.pdf')

SAMPLE_LINES = """
=================================================================================================================================
*     Codotra 202404 202503 202603 sprl  *                ACHATS                  *        18.10.26  02:12        PAGE :    1   *
=================================================================================================================================

  =================================================================================================================
  :   DATE   :  PIECE  :                 COMPTE                  :   DEBIT    :   CREDIT   :   COMMENTAIRE        :
  :----------:---------:-----------------------------------------:------------:------------:----------------------:
  : 30.06.25 :F       1:-----------------------------------------:------------:------------:----------------------:
  :          :         : F 000031 Resto                          :            :    1188,50 : Resto                :
  :          :         : G 613310 FRAIS DE RESTAURANT - BELGIQUE :    1188,50 :            : Resto                :
  : 30.06.25 :F       2:-----------------------------------------:------------:------------:----------------------:
  :          :         : F 000113 Google                         :            :       0,74 : Google               :
=================================================================================================================================
*     Codotra 202404 202503 202603 sprl  *                ACHATS                  *        18.10.26  02:12        PAGE :    2   *
=================================================================================================================================
  =================================================================================================================
  :   DATE   :  PIECE  :                 COMPTE                  :   DEBIT    :   CREDIT   :   COMMENTAIRE        :
  :----------:---------:-----------------------------------------:------------:------------:----------------------:
  :          :         : G 600000 ACHATS                         :       0,74 :            : Note: x              :
  :----------:---------:-----------------------------------------:------------:------------:----------------------:
  :          :         : Nb Doc: 2         Nb lignes: 4          :    1189,24 :    1189,24 :                      :
  =================================================================================================================
   ===========================================================================
   : COMPTE :            LIBELLE             :    DEBIT       :   CREDIT     :
   : 411000 : TVA DEDUCTIBLE                 :        385,41 :               :
""".splitlines()


class TestParseAmount(unittest.TestCase):

    def test_amounts(self):
        self.assertEqual(parse_amount('  74,50 '), Decimal('74.50'))
        self.assertEqual(parse_amount('1.188,50'), Decimal('1188.50'))
        self.assertEqual(parse_amount('-1,85'), Decimal('-1.85'))
        self.assertIsNone(parse_amount('   '))
        self.assertIsNone(parse_amount('abc'))


class TestParseJournalLines(unittest.TestCase):

    def setUp(self):
        self.entries = parse_journal_lines(SAMPLE_LINES)

    def test_only_entry_lines_are_kept(self):
        self.assertEqual(len(self.entries), 4)

    def test_entry_fields(self):
        self.assertEqual(self.entries[0], {
            'Date': '30.06.25',
            'Pièce': 'F 1',
            'Type': 'F',
            'Compte': '000031',
            'Libellé': 'Resto',
            'Débit': None,
            'Crédit': Decimal('1188.50'),
            'Commentaire': 'Resto',
        })
        self.assertEqual(self.entries[1]['Libellé'], 'FRAIS DE RESTAURANT - BELGIQUE')
        self.assertEqual(self.entries[1]['Débit'], Decimal('1188.50'))

    def test_document_continues_across_pages(self):
        last = self.entries[-1]
        self.assertEqual(last['Pièce'], 'F 2')
        self.assertEqual(last['Compte'], '600000')
        self.assertEqual(last['Commentaire'], 'Note: x')


class TestWriteExcel(unittest.TestCase):

    def test_write_and_read_back(self):
        from openpyxl import load_workbook

        entries = parse_journal_lines(SAMPLE_LINES)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'out.xlsx')
            write_excel(entries, path)
            rows = list(load_workbook(path).active.values)

        self.assertEqual(list(rows[0]), COLUMNS)
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[1][0].date(), date(2025, 6, 30))
        self.assertEqual(rows[1][6], 1188.5)
        self.assertIsNone(rows[1][5])


@unittest.skipUnless(os.path.isfile(SAMPLE_PDF), 'sample PDF not available')
class TestSamplePdf(unittest.TestCase):

    def test_convert_sample_pdf(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'out.xlsx')
            self.assertEqual(main([SAMPLE_PDF, '--output', path]), 0)
            rows = list(load_workbook(path).active.values)[1:]

        # Values from the "Nb Doc: 40  Nb lignes: 116" totals line of the report
        self.assertEqual(len(rows), 116)
        self.assertEqual(len({row[1] for row in rows}), 40)
        self.assertAlmostEqual(sum(row[5] or 0 for row in rows), 4677.53, places=2)
        self.assertAlmostEqual(sum(row[6] or 0 for row in rows), 4677.53, places=2)

    def test_missing_file(self):
        self.assertEqual(main(['does-not-exist.pdf']), 1)


if __name__ == '__main__':
    unittest.main()
