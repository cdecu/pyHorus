#!/usr/bin/env python3
"""
Journal PDF to Excel converter

Reads a journal report PDF (e.g. docs/2025.05.FA.pdf) and exports only the
journal entries table:

    :   DATE   :  PIECE  :   COMPTE   :   DEBIT    :   CREDIT   :   COMMENTAIRE   :

to an Excel file. The other sections of the report (VAT summary,
"CENTRALISATION DU JOURNAL", ...) are ignored.

Usage:
    python pdf_to_excel.py docs/2025.05.FA.pdf
    python pdf_to_excel.py docs/2025.05.FA.pdf --output journal.xlsx
"""

import argparse
import os
import re
import sys
from decimal import Decimal, InvalidOperation
from typing import Iterable, List, Optional

COLUMNS = ['Date', 'Pièce', 'Type', 'Compte', 'Libellé', 'Débit', 'Crédit', 'Commentaire']

_DATE_RE = re.compile(r'^\d{2}\.\d{2}\.\d{2,4}$')
_ACCOUNT_RE = re.compile(r'^([A-Z])\s+(\S+)\s*(.*)$')


def parse_amount(value: str) -> Optional[Decimal]:
    """Convert an amount such as '1.188,50' or '74,50' to a Decimal."""
    value = value.strip().replace(' ', '').replace('.', '').replace(',', '.')
    if not value:
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def _split_row(line: str) -> Optional[List[str]]:
    """Split a ':'-delimited table row into its 6 cells, or return None."""
    stripped = line.strip()
    if not (stripped.startswith(':') and stripped.endswith(':')):
        return None
    cells = stripped[1:-1].split(':')
    if len(cells) < 6:
        return None
    if len(cells) > 6:
        # A ':' inside the comment: merge the extra cells back together
        cells = cells[:5] + [':'.join(cells[5:])]
    return [c.strip() for c in cells]


def _is_header(cells: List[str]) -> bool:
    return [c.upper() for c in cells[:2]] == ['DATE', 'PIECE'] and cells[2].upper() == 'COMPTE'


def parse_journal_lines(lines: Iterable[str]) -> List[dict]:
    """
    Parse the text lines of the report and return the journal entries.

    The table is only read between its header line (DATE / PIECE / COMPTE ...)
    and the next '=====' border, so it can span several pages.
    """
    entries = []
    in_table = False
    current_date = ''
    current_piece = ''

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith('='):
            in_table = False
            continue

        cells = _split_row(line)
        if cells is None:
            continue

        if _is_header(cells):
            in_table = True
            continue
        if not in_table:
            continue

        date, piece, account, debit, credit, comment = cells

        # Separator line ":----------:---------:..."
        if all(set(cell) <= {'-'} for cell in cells):
            continue

        # Document line: ": 30.06.25 :F       1:-------...:"
        if _DATE_RE.match(date):
            current_date = date
            current_piece = re.sub(r'\s+', ' ', piece)
            continue

        # Totals line: "Nb Doc: 40  Nb lignes: 116"
        if account.startswith('Nb Doc'):
            continue

        match = _ACCOUNT_RE.match(account)
        if match:
            account_type, account_number, label = match.groups()
        else:
            account_type, account_number, label = '', '', account

        entries.append({
            'Date': current_date,
            'Pièce': current_piece,
            'Type': account_type,
            'Compte': account_number,
            'Libellé': label.strip(),
            'Débit': parse_amount(debit),
            'Crédit': parse_amount(credit),
            'Commentaire': comment,
        })

    return entries


def extract_text_lines(pdf_path: str) -> List[str]:
    """Extract the text of the PDF, preserving the layout of each line."""
    from pypdf import PdfReader

    reader = PdfReader(pdf_path)
    lines = []
    for page in reader.pages:
        text = page.extract_text(extraction_mode='layout') or ''
        lines.extend(text.splitlines())
    return lines


def write_excel(entries: List[dict], output_path: str) -> None:
    """Write the journal entries to an Excel file."""
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Journal'
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for entry in entries:
        row = []
        for column in COLUMNS:
            value = entry[column]
            if column == 'Date' and value:
                fmt = '%d.%m.%y' if len(value) == 8 else '%d.%m.%Y'
                value = datetime.strptime(value, fmt).date()
            elif isinstance(value, Decimal):
                value = float(value)
            row.append(value)
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        row[0].number_format = 'DD/MM/YYYY'
        row[5].number_format = '#,##0.00'
        row[6].number_format = '#,##0.00'

    widths = {'A': 12, 'B': 10, 'C': 6, 'D': 10, 'E': 40, 'F': 12, 'G': 12, 'H': 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    wb.save(output_path)


def create_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description='Convert the journal table of a PDF report to an Excel file',
    )
    parser.add_argument('pdf', help='Path to the PDF report')
    parser.add_argument(
        '--output', '-o',
        help='Path to the Excel file (default: same name as the PDF with .xlsx)',
    )
    return parser


def main(argv: Optional[List[str]] = None) -> int:
    args = create_parser().parse_args(argv)

    if not os.path.isfile(args.pdf):
        print(f"Error: file not found: {args.pdf}", file=sys.stderr)
        return 1

    output = args.output or os.path.splitext(args.pdf)[0] + '.xlsx'

    entries = parse_journal_lines(extract_text_lines(args.pdf))
    if not entries:
        print("Error: no journal entries found in the PDF", file=sys.stderr)
        return 1

    write_excel(entries, output)

    total_debit = sum(e['Débit'] or 0 for e in entries)
    total_credit = sum(e['Crédit'] or 0 for e in entries)
    print(f"{len(entries)} lines written to {output} "
          f"(debit: {total_debit}, credit: {total_credit})")
    return 0


if __name__ == '__main__':
    sys.exit(main())
