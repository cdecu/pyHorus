#!/usr/bin/env python3
"""
Journal PDF to UBL invoices converter

Reads a purchase journal report PDF (e.g. docs/2025.05.FA.pdf) the same way
as pdf_to_excel.py (journal table "DATE : PIECE : COMPTE : DEBIT : CREDIT :
COMMENTAIRE") and writes one UBL 2.1 invoice per document, following the
Peppol BIS Billing 3.0 specification, in the given output directory.

The supplier of each document is the 'F' account of the document; its details
are read from the third parties Excel file (column 'Clé Fournisseur').
The buyer is CODOTRA (BE0445314330) by default.

VAT is derived from the journal lines:
- 411xxx (deductible VAT) + 451xxx (VAT due): reverse charge, the supplier
  invoice has no VAT (category K for EU suppliers, AE otherwise);
- 411xxx only: standard rated (category S, rate deduced from the amounts;
  the part of the base not covered by the VAT is exported as exempt);
- no VAT line: exempt (category E).

Usage:
    python pdf_to_ubls.py docs/2025.05.FA.pdf output_dir
    python pdf_to_ubls.py docs/2025.05.FA.pdf output_dir --tiers docs/Tiers.xlsx
"""

import argparse
import os
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Dict, Iterable, List, Optional

from pdf_to_excel import extract_text_lines, parse_journal_lines

BUYER_NAME = 'CODOTRA'
BUYER_VAT = 'BE0445314330'
CURRENCY = 'EUR'

CUSTOMIZATION_ID = 'urn:cen.eu:en16931:2017#compliant#urn:fdc:peppol.eu:2017:poacc:billing:3.0'
PROFILE_ID = 'urn:fdc:peppol.eu:2017:poacc:billing:01:1.0'

NS = {
    'invoice': 'urn:oasis:names:specification:ubl:schema:xsd:Invoice-2',
    'creditnote': 'urn:oasis:names:specification:ubl:schema:xsd:CreditNote-2',
    'cac': 'urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2',
    'cbc': 'urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2',
}
ET.register_namespace('cac', NS['cac'])
ET.register_namespace('cbc', NS['cbc'])

VAT_DEDUCTIBLE_PREFIX = '411'
VAT_DUE_PREFIX = '451'
STANDARD_RATES = (Decimal('21'), Decimal('12'), Decimal('6'))
TOLERANCE = Decimal('0.02')
CENT = Decimal('0.01')

EU_COUNTRIES = {
    'AT', 'BE', 'BG', 'CY', 'CZ', 'DE', 'DK', 'EE', 'ES', 'FI', 'FR', 'GR', 'HR', 'HU',
    'IE', 'IT', 'LT', 'LU', 'LV', 'MT', 'NL', 'PL', 'PT', 'RO', 'SE', 'SI', 'SK',
}
# Country codes used in the third parties file that are not ISO 3166-1 alpha-2
COUNTRY_ALIASES = {'B': 'BE', 'ALL': 'DE', 'SUE': 'SE', 'USA': 'US', 'UK': 'GB', 'EL': 'GR'}
# Peppol electronic address schemes (EAS) based on the VAT number
VAT_ENDPOINT_SCHEMES = {
    'AT': '9914', 'BG': '9926', 'CY': '9928', 'CZ': '9929', 'DE': '9930', 'EE': '9931',
    'ES': '9920', 'FR': '9957', 'GB': '9932', 'GR': '9933', 'HR': '9934', 'HU': '9910',
    'IE': '9935', 'IT': '0211', 'LT': '9937', 'LU': '9938', 'LV': '9939', 'MT': '9943',
    'NL': '9944', 'PL': '9945', 'PT': '9946', 'RO': '9947', 'SE': '9955', 'SI': '9949',
    'SK': '9950',
}
EXEMPTION_REASONS = {
    'E': (None, 'Exempt from VAT'),
    'AE': ('VATEX-EU-AE', 'Reverse charge'),
    'K': ('VATEX-EU-IC', 'Intra-Community supply'),
}

_VAT_RE = re.compile(r'^([A-Z]{2})[0-9A-Z]{2,}$')


def _cents(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def _fmt(value: Decimal) -> str:
    return str(_cents(value))


def _fmt_rate(rate: Decimal) -> str:
    return str(_cents(rate)) if rate == _cents(rate) else str(rate.normalize())


def _str(value) -> str:
    return '' if value is None else str(value).strip()


def _normalize_key(key) -> str:
    key = _str(key).upper()
    return key.lstrip('0') or key


# --------------------------------------------------------------------------
# Third parties
# --------------------------------------------------------------------------

def normalize_vat(vat) -> str:
    """Return the VAT number without separators, or '' if it is not usable."""
    vat = re.sub(r'[\s.\-]', '', _str(vat)).upper()
    match = _VAT_RE.match(vat)
    if not match or match.group(1) == 'ZZ':
        return ''
    return vat


def normalize_country(country, vat: str = '') -> str:
    """ISO country code, taken from the VAT number when available."""
    if vat:
        return 'GR' if vat[:2] == 'EL' else vat[:2]
    country = _str(country).upper()
    country = COUNTRY_ALIASES.get(country, country)
    return country if len(country) == 2 else 'BE'


def load_suppliers(xlsx_path: str) -> Dict[str, dict]:
    """Read the third parties file and index the suppliers by 'Clé Fournisseur'."""
    from openpyxl import load_workbook

    ws = load_workbook(xlsx_path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [_str(h).lower() for h in next(rows)]

    def col(row, name):
        return row[header.index(name)] if name in header else None

    suppliers = {}
    for row in rows:
        key = _str(col(row, 'clé fournisseur'))
        if not key:
            continue
        vat = normalize_vat(col(row, 'id tva'))
        suppliers[_normalize_key(key)] = {
            'key': key,
            'name': _str(col(row, 'nom')),
            'vat': vat,
            'street': _str(col(row, 'adresse')),
            'postal_code': _str(col(row, 'c.postal')),
            'city': _str(col(row, 'localité')),
            'country': normalize_country(col(row, 'pays'), vat),
            'contact': _str(col(row, 'contact')),
            'phone': _str(col(row, 'téléphone')),
            'email': _str(col(row, 'e-mail fact.')) or _str(col(row, 'e-mail')),
            'bank_account': re.sub(r'\s', '', _str(col(row, 'cpt bancaire'))),
        }
    return suppliers


# --------------------------------------------------------------------------
# Documents
# --------------------------------------------------------------------------

def group_documents(entries: Iterable[dict]) -> List[dict]:
    """Group the journal lines by document (date + piece)."""
    documents = []
    for entry in entries:
        key = (entry['Date'], entry['Pièce'])
        if not documents or documents[-1]['key'] != key:
            documents.append({'key': key, 'date': entry['Date'], 'piece': entry['Pièce'], 'lines': []})
        documents[-1]['lines'].append(entry)
    return documents


def _balance(entry: dict) -> Decimal:
    """Debit - credit of a journal line."""
    return (entry['Débit'] or Decimal(0)) - (entry['Crédit'] or Decimal(0))


def _effective_rate(base: Decimal, vat: Decimal) -> Decimal:
    return (vat * 100 / base).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP) if base else Decimal(0)


def _tax_split(base: Decimal, vat: Decimal, warnings: List[str]) -> List[dict]:
    """Return the VAT subtotals of a standard rated document."""
    for rate in STANDARD_RATES:
        if abs(base * rate / 100 - vat) <= TOLERANCE:
            return [{'category': 'S', 'rate': rate, 'taxable': base, 'tax': vat}]
    if base > 0:
        for rate in STANDARD_RATES:
            taxable = _cents(vat * 100 / rate)
            if taxable <= base:
                subtotals = [{'category': 'S', 'rate': rate, 'taxable': taxable, 'tax': vat}]
                if taxable < base:
                    subtotals.append({'category': 'E', 'rate': Decimal(0), 'taxable': base - taxable,
                                      'tax': Decimal(0)})
                return subtotals
    rate = _effective_rate(base, vat)
    warnings.append(f"VAT {vat} does not match a standard rate for base {base}, using {rate}%")
    return [{'category': 'S', 'rate': rate, 'taxable': base, 'tax': vat}]


def build_invoice_data(document: dict, suppliers: Dict[str, dict]) -> dict:
    """Compute the supplier, lines, VAT subtotals and totals of a document."""
    warnings = []
    supplier_lines = [e for e in document['lines'] if e['Type'] == 'F']
    if not supplier_lines:
        raise ValueError(f"document {document['piece']}: no supplier (F) line")
    supplier_line = supplier_lines[0]

    supplier = suppliers.get(_normalize_key(supplier_line['Compte']))
    if supplier is None:
        warnings.append(f"supplier {supplier_line['Compte']} not found in the third parties file")
        supplier = {
            'key': supplier_line['Compte'], 'name': supplier_line['Libellé'], 'vat': '', 'street': '',
            'postal_code': '', 'city': '', 'country': 'BE', 'contact': '', 'phone': '', 'email': '',
            'bank_account': '',
        }

    total = -sum(_balance(e) for e in supplier_lines)
    vat = Decimal(0)
    vat_due = Decimal(0)
    expenses = []
    for entry in document['lines']:
        if entry['Type'] == 'F':
            continue
        if entry['Compte'].startswith(VAT_DEDUCTIBLE_PREFIX):
            vat += _balance(entry)
        elif entry['Compte'].startswith(VAT_DUE_PREFIX):
            vat_due -= _balance(entry)
        else:
            expenses.append(entry)

    credit_note = total < 0
    sign = -1 if credit_note else 1
    total, vat, vat_due = sign * total, sign * vat, sign * vat_due
    amounts = [sign * _balance(e) for e in expenses]
    base = sum(amounts, Decimal(0))

    if vat_due:
        # Reverse charge: the VAT is self-assessed by the buyer
        country = supplier['country']
        category = 'K' if country in EU_COUNTRIES and country != 'BE' else 'AE'
        if vat_due != vat:
            warnings.append(f"deductible VAT {vat} differs from VAT due {vat_due}")
        vat = Decimal(0)
        subtotals = [{'category': category, 'rate': Decimal(0), 'taxable': base, 'tax': vat}]
    elif vat:
        subtotals = _tax_split(base, vat, warnings) if all(a >= 0 for a in amounts) else \
            [{'category': 'S', 'rate': _effective_rate(base, vat), 'taxable': base, 'tax': vat}]
    else:
        subtotals = [{'category': 'E', 'rate': Decimal(0), 'taxable': base, 'tax': vat}]

    if not supplier['vat']:
        warnings.append(f"supplier {supplier['key']} ({supplier['name']}) has no VAT number")
    if base + vat != total:
        warnings.append(f"total {total} differs from base {base} + VAT {vat}")

    # Distribute the base of the expense lines over the VAT subtotals
    lines = []
    remaining = [s['taxable'] for s in subtotals]
    for entry, amount in zip(expenses, amounts):
        for index, subtotal in enumerate(subtotals):
            part = amount if index == len(subtotals) - 1 else min(amount, remaining[index])
            if part == 0 and len(subtotals) > 1:
                continue
            remaining[index] -= part
            amount -= part
            lines.append({
                'name': entry['Libellé'] or entry['Commentaire'] or entry['Compte'],
                'description': entry['Commentaire'],
                'account': entry['Compte'],
                'amount': part,
                'category': subtotal['category'],
                'rate': subtotal['rate'],
            })

    date = datetime.strptime(document['date'], '%d.%m.%y' if len(document['date']) == 8 else '%d.%m.%Y')
    comments = list(dict.fromkeys(e['Commentaire'] for e in document['lines'] if e['Commentaire']))
    return {
        'id': document['piece'].replace(' ', ''),
        'date': date.date().isoformat(),
        'credit_note': credit_note,
        'note': ' / '.join(comments),
        'supplier': supplier,
        'lines': lines,
        'subtotals': subtotals,
        'base': base,
        'vat': vat,
        'total': total,
        'warnings': warnings,
    }


# --------------------------------------------------------------------------
# UBL
# --------------------------------------------------------------------------

def _el(parent, tag: str, text=None, **attrs):
    prefix, name = tag.split(':')
    element = ET.SubElement(parent, f"{{{NS[prefix]}}}{name}", attrs)
    if text is not None:
        element.text = str(text)
    return element


def _amount(parent, tag: str, value: Decimal):
    return _el(parent, tag, _fmt(value), currencyID=CURRENCY)


def _tax_category(parent, tag: str, category: str, rate: Decimal, with_reason: bool = False):
    element = _el(parent, tag)
    _el(element, 'cbc:ID', category)
    _el(element, 'cbc:Percent', _fmt_rate(rate) if category == 'S' else '0')
    if with_reason and category in EXEMPTION_REASONS:
        code, reason = EXEMPTION_REASONS[category]
        if code:
            _el(element, 'cbc:TaxExemptionReasonCode', code)
        _el(element, 'cbc:TaxExemptionReason', reason)
    _el(_el(element, 'cac:TaxScheme'), 'cbc:ID', 'VAT')


def _party(parent, tag: str, name: str, vat: str, country: str, street: str = '', city: str = '',
           postal_code: str = '', contact: str = '', phone: str = '', email: str = '',
           identifier: str = ''):
    party = _el(_el(parent, tag), 'cac:Party')
    if vat and country == 'BE':
        _el(party, 'cbc:EndpointID', vat[2:], schemeID='0208')
    elif vat and vat[:2] in VAT_ENDPOINT_SCHEMES:
        _el(party, 'cbc:EndpointID', vat, schemeID=VAT_ENDPOINT_SCHEMES[vat[:2]])
    if identifier:
        _el(_el(party, 'cac:PartyIdentification'), 'cbc:ID', identifier)
    _el(_el(party, 'cac:PartyName'), 'cbc:Name', name)
    address = _el(party, 'cac:PostalAddress')
    if street:
        _el(address, 'cbc:StreetName', street)
    if city:
        _el(address, 'cbc:CityName', city)
    if postal_code:
        _el(address, 'cbc:PostalZone', postal_code)
    _el(_el(address, 'cac:Country'), 'cbc:IdentificationCode', country)
    if vat:
        tax_scheme = _el(party, 'cac:PartyTaxScheme')
        _el(tax_scheme, 'cbc:CompanyID', vat)
        _el(_el(tax_scheme, 'cac:TaxScheme'), 'cbc:ID', 'VAT')
    legal = _el(party, 'cac:PartyLegalEntity')
    _el(legal, 'cbc:RegistrationName', name)
    if vat and country == 'BE':
        _el(legal, 'cbc:CompanyID', vat[2:], schemeID='0208')
    if contact or phone or email:
        contact_el = _el(party, 'cac:Contact')
        if contact:
            _el(contact_el, 'cbc:Name', contact)
        if phone:
            _el(contact_el, 'cbc:Telephone', phone)
        if email:
            _el(contact_el, 'cbc:ElectronicMail', email)


def build_ubl(data: dict, buyer_name: str = BUYER_NAME, buyer_vat: str = BUYER_VAT) -> ET.ElementTree:
    """Build the UBL Invoice (or CreditNote) XML tree of a document."""
    credit_note = data['credit_note']
    root_ns = NS['creditnote'] if credit_note else NS['invoice']
    root_name = 'CreditNote' if credit_note else 'Invoice'
    ET.register_namespace('', root_ns)
    root = ET.Element(f"{{{root_ns}}}{root_name}")

    _el(root, 'cbc:CustomizationID', CUSTOMIZATION_ID)
    _el(root, 'cbc:ProfileID', PROFILE_ID)
    _el(root, 'cbc:ID', data['id'])
    _el(root, 'cbc:IssueDate', data['date'])
    _el(root, 'cbc:CreditNoteTypeCode' if credit_note else 'cbc:InvoiceTypeCode',
        '381' if credit_note else '380')
    if data['note']:
        _el(root, 'cbc:Note', data['note'])
    _el(root, 'cbc:DocumentCurrencyCode', CURRENCY)
    _el(root, 'cbc:BuyerReference', data['id'])

    supplier = data['supplier']
    _party(root, 'cac:AccountingSupplierParty', supplier['name'], supplier['vat'], supplier['country'],
           supplier['street'], supplier['city'], supplier['postal_code'],
           supplier['contact'], supplier['phone'], supplier['email'], supplier['key'])
    buyer_vat = normalize_vat(buyer_vat)
    buyer_country = normalize_country(None, buyer_vat)
    _party(root, 'cac:AccountingCustomerParty', buyer_name, buyer_vat, buyer_country)

    if any(s['category'] == 'K' for s in data['subtotals']):
        # Intra-Community supplies require the delivery date and country
        delivery = _el(root, 'cac:Delivery')
        _el(delivery, 'cbc:ActualDeliveryDate', data['date'])
        address = _el(_el(delivery, 'cac:DeliveryLocation'), 'cac:Address')
        _el(_el(address, 'cac:Country'), 'cbc:IdentificationCode', buyer_country)

    if supplier['bank_account']:
        payment = _el(root, 'cac:PaymentMeans')
        _el(payment, 'cbc:PaymentMeansCode', '30')
        _el(_el(payment, 'cac:PayeeFinancialAccount'), 'cbc:ID', supplier['bank_account'])

    tax_total = _el(root, 'cac:TaxTotal')
    _amount(tax_total, 'cbc:TaxAmount', data['vat'])
    for subtotal in data['subtotals']:
        element = _el(tax_total, 'cac:TaxSubtotal')
        _amount(element, 'cbc:TaxableAmount', subtotal['taxable'])
        _amount(element, 'cbc:TaxAmount', subtotal['tax'])
        _tax_category(element, 'cac:TaxCategory', subtotal['category'], subtotal['rate'], with_reason=True)

    totals = _el(root, 'cac:LegalMonetaryTotal')
    _amount(totals, 'cbc:LineExtensionAmount', data['base'])
    _amount(totals, 'cbc:TaxExclusiveAmount', data['base'])
    _amount(totals, 'cbc:TaxInclusiveAmount', data['base'] + data['vat'])
    _amount(totals, 'cbc:PayableAmount', data['total'])

    for number, line in enumerate(data['lines'], start=1):
        element = _el(root, 'cac:CreditNoteLine' if credit_note else 'cac:InvoiceLine')
        _el(element, 'cbc:ID', number)
        _el(element, 'cbc:CreditedQuantity' if credit_note else 'cbc:InvoicedQuantity', '1', unitCode='C62')
        _amount(element, 'cbc:LineExtensionAmount', line['amount'])
        _el(element, 'cbc:AccountingCost', line['account'])
        item = _el(element, 'cac:Item')
        if line['description']:
            _el(item, 'cbc:Description', line['description'])
        _el(item, 'cbc:Name', line['name'])
        _tax_category(item, 'cac:ClassifiedTaxCategory', line['category'], line['rate'])
        _amount(_el(element, 'cac:Price'), 'cbc:PriceAmount', line['amount'])

    tree = ET.ElementTree(root)
    ET.indent(tree)
    return tree


def write_ubls(documents: List[dict], suppliers: Dict[str, dict], output_dir: str,
               buyer_name: str = BUYER_NAME, buyer_vat: str = BUYER_VAT) -> List[str]:
    """Write one UBL file per document and return the paths of the files."""
    os.makedirs(output_dir, exist_ok=True)
    paths = []
    for document in documents:
        data = build_invoice_data(document, suppliers)
        for warning in data['warnings']:
            print(f"Warning: document {document['piece']}: {warning}", file=sys.stderr)
        path = os.path.join(output_dir, f"{data['id']}.xml")
        build_ubl(data, buyer_name, buyer_vat).write(path, encoding='UTF-8', xml_declaration=True)
        paths.append(path)
    return paths


def create_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description='Convert the documents of a PDF journal report to UBL (Peppol BIS 3.0) invoices',
    )
    parser.add_argument('pdf', help='Path to the PDF report')
    parser.add_argument('output_dir', help='Directory where the UBL files are written')
    parser.add_argument(
        '--tiers', '-t',
        help="Third parties Excel file (default: Tiers.xlsx next to the PDF)",
    )
    parser.add_argument('--buyer-name', default=BUYER_NAME, help=f'Buyer name (default: {BUYER_NAME})')
    parser.add_argument('--buyer-vat', default=BUYER_VAT, help=f'Buyer VAT number (default: {BUYER_VAT})')
    return parser


def main(argv: Optional[List[str]] = None) -> int:
    args = create_parser().parse_args(argv)
    tiers = args.tiers or os.path.join(os.path.dirname(args.pdf), 'Tiers.xlsx')

    for path in (args.pdf, tiers):
        if not os.path.isfile(path):
            print(f"Error: file not found: {path}", file=sys.stderr)
            return 1

    documents = group_documents(parse_journal_lines(extract_text_lines(args.pdf)))
    if not documents:
        print("Error: no journal entries found in the PDF", file=sys.stderr)
        return 1

    paths = write_ubls(documents, load_suppliers(tiers), args.output_dir, args.buyer_name, args.buyer_vat)
    print(f"{len(paths)} UBL files written to {args.output_dir}")
    return 0


if __name__ == '__main__':
    sys.exit(main())
